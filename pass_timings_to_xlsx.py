#!/usr/bin/env python3
"""Convert pass-timing NDJSON records to XLSX.

Input is the output of replay_pass_timings.py. Each row is one compilation,
with kernel_signature and the timing breakdown metrics.
"""

from __future__ import annotations

import argparse
import json
import sys

from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("openpyxl is required: python3 -m pip install openpyxl") from exc


METRIC_NAMES = [
    "total_us",
    "npuir_self_us",
    "hivmc_launch_us",
    "hivmc_self_us",
    "bisheng_self_us",
    "lld_self_us",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "inputs",
        nargs="+",
        type=Path,
        help="NDJSON file(s) or directory/directories with *.ndjson files.",
    )
    parser.add_argument(
        "-o",
        "--out",
        required=True,
        type=Path,
        help="Output XLSX path.",
    )
    parser.add_argument(
        "--skip-bad-lines",
        action="store_true",
        help="Warn and skip malformed JSON lines instead of failing.",
    )
    return parser.parse_args()


def input_paths(inputs: list[Path]) -> list[Path]:
    paths: list[Path] = []
    for path in inputs:
        if path.is_dir():
            paths.extend(sorted(path.rglob("*.ndjson")))
        elif path.is_file():
            paths.append(path)
        else:
            raise SystemExit(f"input does not exist: {path}")
    if not paths:
        raise SystemExit("no NDJSON files found")
    return paths


def load_records(
    paths: list[Path], *, skip_bad_lines: bool
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []

    for path in paths:
        with path.open(encoding="utf-8") as handle:
            for line_no, line in enumerate(handle, 1):
                line = line.strip()
                if not line:
                    continue
                try:
                    record = json.loads(line)
                except json.JSONDecodeError as exc:
                    if not skip_bad_lines:
                        raise SystemExit(f"bad JSON {path}:{line_no}: {exc}") from exc
                    print(f"warning: bad JSON {path}:{line_no}: {exc}", file=sys.stderr)
                    continue

                if not isinstance(record, dict):
                    raise SystemExit(f"invalid record at {path}:{line_no}: expected object")

                record["_input_file"] = str(path)
                record["_input_line"] = line_no
                records.append(record)

    if not records:
        raise SystemExit("no records loaded")
    return records


def append_header(ws, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize_columns(ws, max_width: int = 100) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            widths[cell.column] = min(
                max(widths.get(cell.column, 0), len(value) + 2), max_width
            )
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = max(width, 10)


def check_sum(record: dict[str, Any]) -> int | None:
    total = record.get("total_us")
    parts = [record.get(name) for name in METRIC_NAMES[1:]]  # skip total_us
    if not isinstance(total, (int, float)):
        return None
    if any(v is None or not isinstance(v, (int, float)) for v in parts):
        return None
    return int(sum(parts))

def write_records_sheet(wb: Workbook, records: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "pass_timings"

    header = [
        "input_file",
        "input_line",
        "kernel_signature",
        "source",
    ] + METRIC_NAMES + ["sum_check"]

    append_header(ws, header)

    for record in records:
        total = record.get("total_us")
        calc_sum = check_sum(record)

        ws.append(
            [
                record.get("_input_file", ""),
                record.get("_input_line", ""),
                record.get("kernel_signature", ""),
                record.get("source", ""),
            ]
            + [record.get(name, "") for name in METRIC_NAMES]
            + [
                "OK"
                if calc_sum == total
                else f"MISMATCH ({calc_sum} vs {total})"
                if calc_sum is not None
                else "?"
            ]
        )

    autosize_columns(ws)


def main() -> int:
    args = parse_args()
    paths = input_paths(args.inputs)
    records = load_records(paths, skip_bad_lines=args.skip_bad_lines)

    wb = Workbook()
    write_records_sheet(wb, records)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {len(records)} records from {len(paths)} file(s): {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
